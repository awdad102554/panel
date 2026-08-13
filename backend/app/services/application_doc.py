"""仲裁申请书生成:拼接文本 + 模板占位符替换。

文本拼接格式与 InfoWeb static/js/scripts.js 的 buildXxx 逐一对齐;
占位符替换与 InfoWeb routes/application.py 一致(仿宋 15pt、申请人/被申请人前缀加粗、
多行缩进)。模板:templates/仲裁申请书模板.docx(占位符 {applicant_info} 等 5 个)。
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

_TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / "仲裁申请书模板.docx"

_DOC_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _cn_date(value: str | None, with_day: bool = True) -> str:
    """YYYY-MM-DD → YYYY年MM月DD日(或到月);空值返回占位串。"""
    if not value:
        return "XXXX年XX月XX日" if with_day else "XXXX年XX月"
    for fmt in ("%Y-%m-%d", "%Y-%m"):
        try:
            d = datetime.strptime(value[:10], fmt)
            if with_day and fmt == "%Y-%m":
                d = d.replace(day=1)
            return (
                f"{d.year}年{d.month:02d}月{d.day:02d}日"
                if with_day
                else f"{d.year}年{d.month:02d}月"
            )
        except ValueError:
            continue
    return value


def build_applicant_info(applicants: list[dict]) -> str:
    """申请人:姓名,性别,民族,XXXX年XX月XX日出生,身份证住址:XXX。公民身份号码:XXX,电话:XXX。"""
    lines = []
    for index, app in enumerate(applicants):
        seq = str(index + 1) if len(applicants) > 1 else ""
        parts: list[str] = []
        if app.get("name"):
            parts.append(f"申请人{seq}：{app['name']}")
        if app.get("gender"):
            parts.append(f"，{app['gender']}")
        if app.get("ethnicity"):
            parts.append(f"，{app['ethnicity']}")
        if app.get("birth_date"):
            parts.append(f"，{_cn_date(app['birth_date'])}出生")
        if app.get("address"):
            parts.append(f"，身份证住址：{app['address']}")
        if app.get("id_number"):
            parts.append(f"。公民身份号码：{app['id_number']}")
        if app.get("phone"):
            parts.append(f"，电话：{app['phone']}")
        parts.append("。")
        lines.append("".join(parts))
    return "\n".join(lines)


def build_respondent_info(respondents: list[dict]) -> str:
    """被申请人：名称，住所：XXX，统一社会信用代码：XXX。\\n法定代表人：XXX；职务，电话：XXX。

    拼接直接用全角符号,不对用户内容做全局替换(避免误改地址等字段里的半角符号)。
    """
    blocks = []
    for index, resp in enumerate(respondents):
        seq = str(index + 1) if len(respondents) > 1 else ""
        text = f"被申请人{seq}：{resp.get('name') or '未知'}"
        if resp.get("address"):
            text += f"，住所：{resp['address']}"
        if resp.get("id_number"):
            text += f"，统一社会信用代码：{resp['id_number']}"
        text += "。"
        if resp.get("legal_person") or resp.get("legal_title") or resp.get("phone"):
            text += "\n"  # 换行,缩进由替换逻辑统一处理
            if resp.get("legal_person"):
                text += f"法定代表人：{resp['legal_person']}"
            if resp.get("legal_title"):
                text += f"；{resp['legal_title']}"
            if resp.get("phone"):
                text += f"，电话：{resp['phone']}"
            text += "。"
        blocks.append(text)
    return "\n".join(blocks)


def _party_requests(app: dict) -> list[tuple[str, dict | None]]:
    """当事人请求行 [(文本, 结构化spec)]:requests 与 request_specs 按索引对齐;
    旧数据没有 requests 时退回单一 claims 文本(与提交链 build_draft 的兜底一致)。"""
    specs = app.get("request_specs") or []
    rows: list[tuple[str, dict | None]] = []
    for i, req in enumerate(app.get("requests") or []):
        if not req or not req.strip():
            continue
        spec = specs[i] if i < len(specs) and isinstance(specs[i], dict) else None
        rows.append((req.strip(), spec))
    if not rows:
        claims = (app.get("claims") or "").strip()
        if claims:
            rows.append((claims, None))
    return rows


def _disambiguate_applicant(text: str, seq: int) -> str:
    """多申请人案件:请求文本里泛指的"申请人"→"申请人N"(seq 从 1 起)。
    与 case_submit 同名函数同一正则;"被申请人"/已编号的"申请人N"不受影响。"""
    return re.sub(r"(?<!被)申请人(?!\d)", f"申请人{seq}", text)


def build_requests_text(applicants: list[dict], respondents: list[dict] | None = None) -> str:
    """所有申请人的请求合并编号;已带数字序号的不重复添加。
    多申请人时给每条请求的"申请人"加编号,多被申请人时支付类表述改
    "各被申请人共同",与内网提交口径一致。"""
    multi = len([a for a in applicants if a.get("name")]) > 1
    plural_resp = len([r for r in (respondents or []) if r.get("name")]) > 1
    all_requests: list[str] = []
    request_index = 1
    for ai, app in enumerate(applicants):
        for req, _ in _party_requests(app):
            if multi:
                req = _disambiguate_applicant(req, ai + 1)
            if plural_resp:
                req = req.replace("被申请人向申请人", "各被申请人共同向申请人")
            if re.match(r"^\d+[\.、]", req):
                all_requests.append(req)
            else:
                all_requests.append(f"{request_index}.{req}")
            request_index += 1
    return "\n".join(all_requests)


def build_total_amount(applicants: list[dict]) -> str:
    """两个及以上请求时求和:以上共计XXX.XX元。

    金额取值优先级(与提交链 build_payload 一致):object → items 齐全则求和 →
    文本正则(旧数据兼容)。items 只有部分金额时不求部分和(申请书是正式文书,
    部分和会偏低误导;提交链对缺金额是 422 拒绝)。"""
    rows = [row for app in applicants for row in _party_requests(app)]
    if len(rows) <= 1:
        return ""
    total = 0.0
    has_amount = False
    for text, spec in rows:
        amount: float | None = None
        if spec is not None:
            if spec.get("object") is not None:
                amount = float(spec["object"])
            else:
                items = [i for i in spec.get("items") or [] if isinstance(i, dict)]
                if items and all(i.get("object") is not None for i in items):
                    amount = sum(i["object"] for i in items)
        if amount is None:
            content = text.split("（")[0]
            match = re.search(r"(\d{1,3}(?:,\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?)\s*(万)?元", content)
            if match:
                amount = float(match.group(1).replace(",", ""))
                if match.group(2):
                    amount *= 10000
        if amount:
            total += amount
            has_amount = True
    if has_amount and total > 0:
        return f"以上共计{total:.2f}元。"
    return ""


def build_facts_reasons(applicants: list[dict], respondents: list[dict]) -> str:
    """每人一段(compose_party_facts 拼接,与提交内网同一规则);末尾加通用结尾。"""
    from app.services.case_submit import compose_party_facts  # 延迟导入避免模块级耦合

    parts: list[str] = []
    default_workplace = respondents[0].get("name") if respondents else ""
    total = len(applicants)
    for index, app in enumerate(applicants):
        q = dict(app)
        # 工作地点缺省时沿用 InfoWeb 旧行为:用第一被申请人名称兜底。
        if not q.get("work_place"):
            q["work_place"] = default_workplace
        para = compose_party_facts(q, index + 1, total)
        if para.strip():
            parts.append(para)
    parts.append(
        "申请人为被申请人提供劳动，被申请人依法应当按照劳动合同约定和国家规定，"
        "及时足额按月向申请人支付劳动报酬。现申请人为了维护自身的合法权益，特依法提起仲裁，望予以公正裁决。"
    )
    return "\n".join(parts)


def _fill_paragraph(para, key: str, value: str) -> None:
    """按 InfoWeb application.py 的规则替换占位符:仿宋 15pt、
    申请人/被申请人前缀加粗、多行从第二行起加全角缩进。"""
    para.clear()
    lines = str(value).split("\n")
    for i, line in enumerate(lines):
        if i > 0 and key in ("applicant_info", "respondent_info", "requests", "facts_reasons"):
            indent_run = para.add_run("　　")
            indent_run.font.name = "仿宋"
            indent_run.font.size = Pt(15)
            indent_run.font.bold = False
            indent_run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

        if key in ("applicant_info", "respondent_info"):
            colon_idx = line.find("：")
            if colon_idx == -1:
                colon_idx = line.find(":")
            if colon_idx > 0:
                prefix = line[: colon_idx + 1]
                content = line[colon_idx + 1 :]
                # 申请人X:/被申请人X: 加粗;法定代表人: 不加粗
                is_bold = "申请人" in prefix and "法定代表人" not in prefix
                run = para.add_run(prefix)
                run.font.name = "仿宋"
                run.font.size = Pt(15)
                run.font.bold = is_bold
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
                if content:
                    run = para.add_run(content)
                    run.font.name = "仿宋"
                    run.font.size = Pt(15)
                    run.font.bold = False
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
            else:
                run = para.add_run(line)
                run.font.name = "仿宋"
                run.font.size = Pt(15)
                run.font.bold = False
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
        else:
            run = para.add_run(line)
            run.font.name = "仿宋"
            run.font.size = Pt(15)
            run.font.bold = False
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

        if i < len(lines) - 1:
            para.add_run("\n")


def generate_application(case_data: dict, output_path: Path) -> Path:
    """按案件数据生成申请书 docx,返回输出路径。

    case_data: {"applicants": [...], "respondents": [...]}(预登记案件的当事人 JSON)。
    """
    data = {
        "applicant_info": build_applicant_info(case_data.get("applicants") or []),
        "respondent_info": build_respondent_info(case_data.get("respondents") or []),
        "requests": build_requests_text(
            case_data.get("applicants") or [], case_data.get("respondents") or []
        ),
        "total_amount": build_total_amount(case_data.get("applicants") or []),
        "facts_reasons": build_facts_reasons(
            case_data.get("applicants") or [], case_data.get("respondents") or []
        ),
    }

    doc = Document(_TEMPLATE_PATH)
    for para in doc.paragraphs:
        for key, value in data.items():
            if "{" + key + "}" in para.text:
                _fill_paragraph(para, key, value)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for key, value in data.items():
                        placeholder = "{" + key + "}"
                        if placeholder in para.text:
                            para.text = para.text.replace(placeholder, str(value))
                            for run in para.runs:
                                run.font.name = "仿宋"
                                run.font.size = Pt(15)
                                run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
